import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

SCREENSHOT_DIR = r"d:\ITC-SV\laragon\www\NguyenNgocHaThao_501250482\screenshots"
OUTPUT_PPTX = r"d:\ITC-SV\laragon\www\NguyenNgocHaThao_501250482\BaoCao_ChucNang_CineTicket.pptx"

# Color Palette: Luxury Cinema VIP
BG_COLOR = RGBColor(11, 14, 23)        # #0b0e17
CARD_BG = RGBColor(18, 23, 37)         # #121725
GOLD_ACCENT = RGBColor(243, 207, 122)  # #f3cf7a
GOLD_DARK = RGBColor(212, 175, 55)     # #d4af37
EMERALD = RGBColor(16, 185, 129)       # #10b981
CYAN = RGBColor(56, 189, 248)          # #38bdf8
WHITE = RGBColor(255, 255, 255)
GRAY_TEXT = RGBColor(180, 190, 205)
BORDER_COLOR = RGBColor(40, 50, 75)

def apply_background(slide):
    # Background shape covering full slide
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()
    return bg

def add_header(slide, slide_num, category, title):
    # Category tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.35))
    tf = tag_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"CINETICKET VIP CINEMA • CHỨC NĂNG #{slide_num:02d} • {category.upper()}"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.5), Inches(0.65))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = WHITE

def add_feature_card(slide, bullets, summary_note=None):
    # Left Card background
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.4), Inches(5.4))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_COLOR
    card.line.width = Pt(1.5)

    # Card Content Textbox
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.65), Inches(5.0), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "CHI TIẾT CHỨC NĂNG & ĐIỂM NỔI BẬT"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = GOLD_ACCENT
    p0.space_after = Pt(12)

    for i, bullet in enumerate(bullets):
        p = tf.add_paragraph()
        p.text = f"•  {bullet}"
        p.font.size = Pt(11.5)
        p.font.color.rgb = GRAY_TEXT
        p.space_after = Pt(7)

    if summary_note:
        p_note = tf.add_paragraph()
        p_note.text = f"★ Lưu ý: {summary_note}"
        p_note.font.size = Pt(11)
        p_note.font.bold = True
        p_note.font.color.rgb = EMERALD
        p_note.space_before = Pt(10)

def add_screenshot(slide, image_filename, caption):
    img_path = os.path.join(SCREENSHOT_DIR, image_filename)
    if os.path.exists(img_path):
        # Frame background
        frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.5), Inches(1.5), Inches(6.05), Inches(5.4))
        frame.fill.solid()
        frame.fill.fore_color.rgb = CARD_BG
        frame.line.color.rgb = GOLD_DARK
        frame.line.width = Pt(1.5)

        # Image (width 5.75, height ~ 4.4)
        slide.shapes.add_picture(img_path, Inches(6.65), Inches(1.65), width=Inches(5.75))

        # Caption
        cap_box = slide.shapes.add_textbox(Inches(6.65), Inches(6.35), Inches(5.75), Inches(0.45))
        tf = cap_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"Minh chứng: {caption}"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = GOLD_ACCENT
    else:
        print(f"Warning: Image not found: {img_path}")

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # ==================== SLIDE 1: COVER SLIDE ====================
    s1 = prs.slides.add_slide(blank_layout)
    apply_background(s1)

    # Decorative Card
    cover_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.0), Inches(10.933), Inches(5.5))
    cover_card.fill.solid()
    cover_card.fill.fore_color.rgb = CARD_BG
    cover_card.line.color.rgb = GOLD_DARK
    cover_card.line.width = Pt(2)

    tb = s1.shapes.add_textbox(Inches(1.6), Inches(1.3), Inches(10.133), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "BÁO CÁO TỔNG QUAN HỆ THỐNG & CHỨC NĂNG"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT
    p.space_after = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "CINETICKET VIP CINEMA"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_after = Pt(6)

    p3 = tf.add_paragraph()
    p3.text = "Hệ Thống Đặt Vé Xem Phim Trực Tuyến & Cổng Thanh Toán Mã QR Thông Minh"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.size = Pt(16)
    p3.font.color.rgb = CYAN
    p3.space_after = Pt(28)

    p4 = tf.add_paragraph()
    p4.text = "• Sinh viên thực hiện: Nguyễn Ngọc Hà Thảo  |  MSSV: 501250482"
    p4.alignment = PP_ALIGN.CENTER
    p4.font.size = Pt(14)
    p4.font.bold = True
    p4.font.color.rgb = WHITE
    p4.space_after = Pt(8)

    p5 = tf.add_paragraph()
    p5.text = "• Trụ sở rạp: 12 Trịnh Đình Thảo, Tân Phú, TP.HCM  |  Hotline: 0337191740 (24/7)"
    p5.alignment = PP_ALIGN.CENTER
    p5.font.size = Pt(13)
    p5.font.color.rgb = GRAY_TEXT
    p5.space_after = Pt(6)

    p6 = tf.add_paragraph()
    p6.text = "• Email Admin quản trị hệ thống: hthaongn@gmail.com"
    p6.alignment = PP_ALIGN.CENTER
    p6.font.size = Pt(13)
    p6.font.bold = True
    p6.font.color.rgb = GOLD_ACCENT

    # ==================== SLIDE 2: TRANG CHỦ & HERO SLIDER ====================
    s2 = prs.slides.add_slide(blank_layout)
    apply_background(s2)
    add_header(s2, 1, "Trang Chủ & Giao Diện", "TRANG CHỦ (HOME PAGE) & BANNER BOM TẤN NỔI BẬT")
    add_feature_card(s2, [
        "Banner Hero Slider: Trình chiếu các bom tấn điện ảnh nổi bật với đồ họa chuyển động cao cấp.",
        "Phân loại danh mục phim: Tab 'Phim Đang Chiếu' (Now Showing) và 'Phim Sắp Chiếu' (Coming Soon).",
        "Bộ lọc thể loại tức thì: Lọc nhanh theo Hành động, Khoa học viễn tưởng, Hoạt hình, Tâm lý.",
        "Bảng xếp hạng Top Trending: Vinh danh các bộ phim có lượng vé đặt kỷ lục trong tuần.",
        "Header & Footer Dark Luxury: Tích hợp menu điều hướng thông minh, giỏ vé, liên hệ và chính sách rạp."
    ], "Giao diện chuẩn Dark Luxury hiện đại, tối ưu hiển thị mượt mà trên mọi thiết bị.")
    add_screenshot(s2, "01_trang_chu.png", "Giao diện Trang chủ và Hệ thống rạp CineTicket VIP")

    # ==================== SLIDE 3: DANH MỤC PHIM & BỘ LỌC THỂ LOẠI ====================
    s3 = prs.slides.add_slide(blank_layout)
    apply_background(s3)
    add_header(s3, 2, "Khám Phá Điện Ảnh", "DANH MỤC LỊCH CHIẾU PHIM & BỘ LỌC THỂ LOẠI (/movies)")
    add_feature_card(s3, [
        "Trang danh mục phim toàn diện: Trình bày toàn bộ kho phim chiếu rạp kèm poster và thông số kỹ thuật.",
        "Hệ thống lọc theo thể loại: Các nút bấm lọc nhanh Tất Cả, Hành Động, Khoa Học Viễn Tưởng, Hoạt Hình, Tâm Lý...",
        "Thanh tìm kiếm phim thời gian thực: Tìm kiếm tức thì theo tên phim hoặc đạo diễn.",
        "Thẻ thông tin phim chi tiết: Hiển thị điểm đánh giá sao (IMDb), thời lượng, độ tuổi (T13, T16, T18) và giá vé chuẩn.",
        "Nút Đặt Vé & Xem Chi Tiết: Chuyển hướng trực tiếp sang trang chi tiết phim hoặc luồng đặt vé nhanh chóng."
    ], "Người dùng dễ dàng tra cứu và chọn lọc bộ phim phù hợp nhất chỉ với một cú nhấp chuột.")
    add_screenshot(s3, "02_danh_sach_phim.png", "Trang danh mục phim kèm thanh tìm kiếm và bộ lọc thể loại")

    # ==================== SLIDE 4: CHI TIẾT PHIM & CHỌN SUẤT CHIẾU ====================
    s4 = prs.slides.add_slide(blank_layout)
    apply_background(s4)
    add_header(s4, 3, "Lịch Chiếu Phim", "CHI TIẾT PHIM, TRAILER & KHUNG CHỌN SUẤT CHIẾU")
    add_feature_card(s4, [
        "Thông tin phim toàn diện: Đạo diễn, diễn viên, thời lượng, quốc gia, tóm tắt nội dung cốt truyện kịch tính.",
        "Trailer phim nhúng sắc nét: Khách hàng có thể xem trailer Full HD có phụ đề trực tiếp trên trang.",
        "Hệ thống chọn lịch chiếu thông minh: Lựa chọn cụm rạp, ngày chiếu và các khung giờ chiếu linh hoạt.",
        "Phân phòng chiếu chuyên biệt: Hiển thị rõ số phòng chiếu (Phòng VIP Gold Class, Laser IMAX, Standard).",
        "Đánh giá & Bình luận của khán giả: Xem các phản hồi thực tế và gửi bình luận chấm điểm sao cho phim."
    ], "Khung giờ và phòng chiếu được kết nối trực tiếp với sơ đồ chọn ghế thời gian thực.")
    add_screenshot(s4, "03_chi_tiet_phim.png", "Trang chi tiết phim với trailer, mô tả và khung chọn suất chiếu")

    # ==================== SLIDE 5: SƠ ĐỒ CHỌN GHẾ ====================
    s5 = prs.slides.add_slide(blank_layout)
    apply_background(s5)
    add_header(s5, 4, "Đặt Chỗ Trực Quan", "SƠ ĐỒ CHỌN GHẾ VIP 5 SAO THỜI GIAN THỰC (/seat)")
    add_feature_card(s5, [
        "Mô phỏng rạp chiếu chuẩn quốc tế: Thiết kế màn hình cong trung tâm Cinema Screen với góc nhìn trực quan.",
        "Phân loại vị trí ghế ngồi: Ghế Thường (Standard), Ghế VIP (Gold Class), và Ghế Đôi (Sweetbox Couple).",
        "Trạng thái ghế phân biệt bằng màu sắc: Ghế trống, Ghế đang chọn (Vàng gold), Ghế đã có người đặt (Xám).",
        "Tự động tính tiền: Giá vé được cập nhật tức thì theo thời gian thực dựa trên số lượng và loại ghế đã bấm chọn.",
        "Cơ chế khóa ghế thông minh: Ngăn chặn triệt để tình trạng trùng lặp ghế khi khách hàng đặt chỗ."
    ], "Khách hàng có thể chọn nhiều ghế cùng lúc và hủy chọn linh hoạt trước khi qua bước tiếp theo.")
    add_screenshot(s5, "04_so_do_chon_ghe.png", "Sơ đồ phòng chiếu VIP với màn hình cong và trạng thái ghế thời gian thực")

    # ==================== SLIDE 6: COMBO BẮP NƯỚC ====================
    s6 = prs.slides.add_slide(blank_layout)
    apply_background(s6)
    add_header(s6, 5, "Dịch Vụ Đi Kèm", "THỰC ĐƠN COMBO BẮP NƯỚC & TẠM TÍNH HÓA ĐƠN (/snack)")
    add_feature_card(s6, [
        "Menu ẩm thực rạp phong phú: Bắp rang bơ phô mai, bắp caramel, nước ngọt có ga, nước suối cao cấp.",
        "Combo ưu đãi tiết kiệm: Combo Solo VIP, Combo Couple VIP, Combo Family cho nhóm bạn và gia đình.",
        "Tùy chỉnh số lượng linh hoạt: Nút tăng/giảm (+/-) số lượng món ăn và combo nhanh chóng, trực quan.",
        "Tạm tính hóa đơn thời gian thực: Tự động cộng dồn tiền vé ghế + tiền bắp nước sang bước thanh toán.",
        "Bước đệm tiện lợi: Có thể chọn thêm bắp nước hoặc bấm Tiếp Tục sang Cổng thanh toán."
    ], "Tất cả combo bắp nước đã chọn sẽ được hiển thị chi tiết trong hóa đơn và vé điện tử.")
    add_screenshot(s6, "05_chon_bap_nuoc.png", "Giao diện chọn combo bắp nước và hóa đơn tạm tính")

    # ==================== SLIDE 7: GIỎ VÉ CỦA BẠN (CART) ====================
    s7 = prs.slides.add_slide(blank_layout)
    apply_background(s7)
    add_header(s7, 6, "Quản Lý Đơn Hàng", "GIỎ VÉ CỦA BẠN & QUẢN LÝ VÉ ĐÃ CHỌN (/cart)")
    add_feature_card(s7, [
        "Bảng giỏ vé trực quan: Liệt kê chi tiết danh sách vé phim đã chọn kèm ảnh poster, thể loại và đơn giá.",
        "Tùy chỉnh số lượng vé: Nút tăng/giảm số lượng hoặc xóa vé khỏi giỏ hàng dễ dàng.",
        "Áp dụng mã giảm giá (Voucher): Nhập mã ưu đãi (VIP25, CINETICKET) để giảm trực tiếp 20% trên tổng bill.",
        "Tóm tắt thanh toán rõ ràng: Tạm tính, tiền giảm giá và tổng tiền phải trả được hiển thị minh bạch.",
        "Điều hướng liền mạch: Nút 'Chọn ghế cho phim này' hoặc bấm nút Thanh Toán Chuyển Khoản nhanh chóng."
    ], "Giỏ vé được lưu trữ tự động, đảm bảo khách hàng không bị mất thông tin khi duyệt web.")
    add_screenshot(s7, "06_gio_hang.png", "Giao diện Giỏ vé với danh sách vé phim đã chọn và bảng tóm tắt hóa đơn")

    # ==================== SLIDE 8: CỔNG THANH TOÁN VIETQR ADMIN ====================
    s8 = prs.slides.add_slide(blank_layout)
    apply_background(s8)
    add_header(s8, 7, "Cổng Thanh Toán", "CỔNG THANH TOÁN QUÉT MÃ VIETQR ADMIN MB BANK (/payment)")
    add_feature_card(s8, [
        "Mã VietQR tự động NAPAS: Tạo mã QR ngân hàng chuẩn xác, tự động điền đúng số tiền và cú pháp chuyển khoản.",
        "Thông tin thụ hưởng Admin rõ ràng: Ngân hàng MB Bank, STK 098765432199, Tên chủ TK: Quản Trị Viên VIP.",
        "Tính năng 1-Click Copy thông minh: Sao chép nhanh STK, số tiền và nội dung chuyển khoản chỉ với 1 chạm.",
        "Hỗ trợ đa phương thức: Quét VietQR, Ví MoMo (0337191740), Cổng VNPAY và Thẻ quốc tế Visa/Mastercard.",
        "Khung thông tin người nhận vé: Điền Họ tên, Số điện thoại và Email để nhận vé điện tử ngay sau khi thanh toán."
    ], "Mã QR được hiển thị to rõ ở trung tâm, khách hàng có thể quét bằng bất kỳ app ngân hàng nào.")
    add_screenshot(s8, "07_thanh_toan_vietqr.png", "Cổng thanh toán VietQR hiển thị mã QR và thông tin tài khoản MB Bank")

    # ==================== SLIDE 9: VÉ ĐIỆN TỬ VIP & IN VÉ ====================
    s9 = prs.slides.add_slide(blank_layout)
    apply_background(s9)
    add_header(s9, 8, "Vé Điện Tử", "VÉ ĐIỆN TỬ VIP PASS & TÍNH NĂNG IN VÉ (/order-success)")
    add_feature_card(s9, [
        "Thẻ vé VIP Pass sang trọng: Thiết kế vương giả viền vàng kim với huy hiệu VIP PASS ĐÃ THANH TOÁN.",
        "Mã QR Check-in thời gian thực: Chứa đầy đủ dữ liệu phòng chiếu, ghế ngồi, suất chiếu để nhân viên rạp quét vé.",
        "Thông số vé nổi bật: Hiển thị rõ ràng Cụm rạp, Số phòng chiếu Laser IMAX, Giờ chiếu, Ngày chiếu và Số ghế.",
        "Danh mục bắp nước đính kèm: Liệt kê đầy đủ các món bắp và nước ngọt đã mua trong đơn hàng.",
        "Tính năng In vé điện tử (Print Ticket): Tích hợp hàm in chuẩn trình duyệt (window.print()) hỗ trợ lưu PDF hoặc in giấy."
    ], "Vé điện tử có đầy đủ mã xác thực, mã QR và hóa đơn thanh toán hợp lệ của CineTicket VIP.")
    add_screenshot(s9, "08_ve_dien_tu_thanh_cong.png", "Vé điện tử VIP Pass với mã QR check-in và tính năng in vé")

    # ==================== SLIDE 10: LỊCH SỬ ĐẶT VÉ (MY ORDERS) ====================
    s10 = prs.slides.add_slide(blank_layout)
    apply_background(s10)
    add_header(s10, 9, "Quản Lý Đơn Hàng", "LỊCH SỬ ĐẶT VÉ CỦA TÔI (MY ORDERS) & BỘ LỌC TRẠNG THÁI (/myorders)")
    add_feature_card(s10, [
        "Lưu trữ vé thông minh: Quản lý toàn bộ vé điện tử mà thành viên đã đặt theo trình tự thời gian.",
        "Bộ lọc trạng thái linh hoạt: Lọc vé theo 'Tất cả đơn hàng', 'Đã thanh toán', 'Chờ xác nhận', 'Đã hủy'.",
        "Thẻ vé tóm tắt đầy đủ: Hiển thị tên phim, giờ chiếu, số ghế ngồi, tổng tiền thanh toán và phương thức.",
        "Xem lại mã QR & Chi tiết vé: Cho phép khách hàng mở lại mã QR check-in hoặc xem vé điện tử bất cứ lúc nào.",
        "Hỗ trợ hủy vé / kiểm tra: Khách hàng chủ động theo dõi tình trạng đơn đặt vé trực tiếp trên tài khoản."
    ], "Dữ liệu được bảo lưu tự động giúp khách hàng không bao giờ bị thất lạc vé đã mua.")
    add_screenshot(s10, "09_lich_su_dat_ve.png", "Trang Lịch sử đặt vé với danh sách đơn hàng và bộ lọc trạng thái")

    # ==================== SLIDE 11: QUẢN LÝ TÀI KHOẢN CÁ NHÂN ====================
    s11 = prs.slides.add_slide(blank_layout)
    apply_background(s11)
    add_header(s11, 10, "Tài Khoản Thành Viên", "QUẢN LÝ THÔNG TIN TÀI KHOẢN CÁ NHÂN & BẢO MẬT (/account)")
    add_feature_card(s11, [
        "Hồ sơ thành viên VIP: Hiển thị Họ và tên (Nguyễn Ngọc Hà Thảo), Email, Số điện thoại và cấp bậc VIP Member.",
        "Cập nhật thông tin cá nhân: Chỉnh sửa họ tên, email, số điện thoại và bấm 'Lưu Thay Đổi'.",
        "Đổi mật khẩu tài khoản: Cơ chế đổi mật khẩu an toàn với xác nhận mật khẩu cũ và mật khẩu mới.",
        "Đồng bộ lưu trữ đa tầng: Lưu trữ tức thì vào cả LocalStorage và SessionStorage, không bị mất khi F5 hoặc mở lại web.",
        "Bảo mật & Đăng xuất an toàn: Nút đăng xuất xóa phiên làm việc an toàn trên trình duyệt."
    ], "Toàn bộ thông tin cá nhân được mã hóa và lưu trữ nhất quán trên toàn hệ thống.")
    add_screenshot(s11, "10_quan_ly_tai_khoan.png", "Trang Quản lý tài khoản cá nhân, cập nhật hồ sơ và đổi mật khẩu")

    # ==================== SLIDE 12: ĐĂNG NHẬP & ĐĂNG KÝ THÀNH VIÊN ====================
    s12 = prs.slides.add_slide(blank_layout)
    apply_background(s12)
    add_header(s12, 11, "Xác Thực Người Dùng", "ĐĂNG NHẬP & ĐĂNG KÝ TÀI KHOẢN THÀNH VIÊN VIP (/login)")
    add_feature_card(s12, [
        "Giao diện đăng nhập chuẩn Dark Luxury: Thiết kế hiệu ứng kính thủy tinh Glassmorphism sang trọng và cao cấp.",
        "Phân quyền truy cập 2 vai trò: Thành viên xem phim (Member) và Quản trị viên hệ thống (Admin).",
        "Nút điền nhanh Demo tiện lợi: 1-Click để đăng nhập nhanh tài khoản Khách Hàng VIP hoặc Admin Quản Trị.",
        "Form đăng ký tài khoản mới: Tích hợp xác thực dữ liệu đầu vào (Validation) đầy đủ họ tên, email, mật khẩu.",
        "Lưu trữ phiên đăng nhập: Ghi nhớ thông tin phiên người dùng và bảo mật quyền truy cập các trang nội bộ."
    ], "Xác thực an toàn, bảo vệ tài khoản và phân quyền truy cập chặt chẽ giữa User và Admin.")
    add_screenshot(s12, "11_dang_nhap.png", "Giao diện Đăng nhập với các nút hỗ trợ đăng nhập nhanh tài khoản VIP")

    # ==================== SLIDE 13: ADMIN DASHBOARD - THỐNG KÊ KPI ====================
    s13 = prs.slides.add_slide(blank_layout)
    apply_background(s13)
    add_header(s13, 12, "Bảng Điều Khiển Admin", "ADMIN DASHBOARD: THỐNG KÊ DOANH THU & CHỈ SỐ KPI VẬN HÀNH (/admin)")
    add_feature_card(s13, [
        "Thống kê tổng quan KPI 4 chỉ số: Tổng doanh thu (VNĐ), Tổng số vé đã bán, Số phim đang chiếu, Khách hàng thành viên.",
        "Biểu đồ trực quan hóa dữ liệu: Theo dõi biến động doanh thu theo tuần/tháng và tỷ lệ lấp đầy ghế các phòng chiếu.",
        "Bảng hoạt động giao dịch gần đây: Cập nhật các đơn đặt vé mới nhất kèm trạng thái thanh toán thời gian thực.",
        "Phân quyền quản trị nghiêm ngặt: Chỉ tài khoản có vai trò 'admin' mới được phép truy cập và xem dữ liệu.",
        "Giao diện Dashboard phẳng: Tối ưu tốc độ tải trang và hỗ trợ thao tác điều hành nhanh cho ban quản lý rạp."
    ], "Giao diện quản trị phong cách Dashboard công nghệ cao, cập nhật số liệu tự động.")
    add_screenshot(s13, "12_admin_thong_ke.png", "Bảng điều khiển Quản trị viên với các chỉ số KPI doanh thu và biểu đồ")

    # ==================== SLIDE 14: ADMIN - QUẢN LÝ PHIM ====================
    s14 = prs.slides.add_slide(blank_layout)
    apply_background(s14)
    add_header(s14, 13, "Quản Trị Nội Dung", "ADMIN DASHBOARD: QUẢN LÝ KHO PHIM & SUẤT CHIẾU RẠP")
    add_feature_card(s14, [
        "Bảng danh sách phim chiếu rạp: Hiển thị đầy đủ ảnh poster, tên phim, thể loại, thời lượng, giá vé và trạng thái.",
        "Thêm phim mới (Add Movie): Modal nhập liệu phim mới với poster URL, trailer, tóm tắt và phân loại độ tuổi.",
        "Chỉnh sửa thông tin phim (Edit Movie): Cập nhật giá vé, thể loại, đạo diễn và diễn viên nhanh chóng.",
        "Xóa phim khỏi lịch chiếu: Cho phép xóa hoặc ẩn các bộ phim đã ngừng chiếu tại rạp khỏi hệ thống.",
        "Đồng bộ tức thì: Dữ liệu phim cập nhật sẽ phản ánh ngay lập tức trên trang chủ và danh mục cho khách hàng."
    ], "Toàn bộ kho phim được quản trị tập trung, đảm bảo tính cập nhật liên tục cho rạp.")
    add_screenshot(s14, "13_admin_quan_ly_phim.png", "Tab Quản lý phim với danh sách phim chiếu rạp và các nút Thêm/Sửa/Xóa")

    # ==================== SLIDE 15: ADMIN - CÀI ĐẶT MÃ QR THANH TOÁN ====================
    s15 = prs.slides.add_slide(blank_layout)
    apply_background(s15)
    add_header(s15, 14, "Quản Trị Thanh Toán", "ADMIN DASHBOARD: CÀI ĐẶT MÃ QR & TÀI KHOẢN NGÂN HÀNG THỤ HƯỞNG")
    add_feature_card(s15, [
        "Cấu hình tài khoản ngân hàng nhận tiền: Tùy chọn ngân hàng (MB Bank, VCB, Vietin...), Số tài khoản, Tên chủ tài khoản.",
        "Cấu hình ví MoMo: Nhập số điện thoại ví MoMo Admin (0337191740) để nhận thanh toán tức thời.",
        "Live VietQR Preview: Khung mô phỏng hiển thị trực tiếp mã VietQR mà khách hàng sẽ nhìn thấy khi thanh toán.",
        "Lưu cấu hình thời gian thực: Toàn bộ thay đổi về STK và thông tin ngân hàng được lưu bền vững vào hệ thống.",
        "Đảm bảo đối soát chuẩn xác: Giúp dòng tiền từ vé xem phim đổ trực tiếp về tài khoản chính chủ của Admin."
    ], "Tính năng cốt lõi giúp hệ thống rạp linh hoạt thay đổi tài khoản ngân hàng thụ hưởng bất kỳ lúc nào.")
    add_screenshot(s15, "14_admin_cai_dat_qr.png", "Tab Cài đặt mã QR Admin với công cụ cấu hình tài khoản và live preview")

    # ==================== SLIDE 16: LIÊN HỆ & TỔNG KẾT ====================
    s16 = prs.slides.add_slide(blank_layout)
    apply_background(s16)
    add_header(s16, 15, "Tổng Kết & Liên Hệ", "TRANG LIÊN HỆ, HOTLINE 24/7 & BẢN ĐỒ VỊ TRÍ TRỤ SỞ RẠP (/contact)")
    add_feature_card(s16, [
        "Địa chỉ trụ sở rạp chiếu: 12 Trịnh Đình Thảo, Phường Hòa Thạnh, Quận Tân Phú, TP. Hồ Chí Minh.",
        "Hotline chăm sóc khách hàng: 0337191740 (Hỗ trợ trực tuyến 24/7 giải đáp mọi thắc mắc về vé và suất chiếu).",
        "Email liên hệ chính thức: hthaongn@gmail.com (Tiếp nhận phản hồi và đối soát đơn hàng).",
        "Tích hợp bản đồ trực quan: Khung bản đồ Google Maps định vị cụm rạp CineTicket VIP giúp khách hàng dễ dàng tìm đường.",
        "Form gửi liên hệ trực tuyến: Khách hàng có thể để lại lời nhắn, góp ý trực tiếp trên trang web."
    ], "Hệ thống đáp ứng 100% yêu cầu đồ án với đầy đủ tính năng thực tế và dữ liệu đồng bộ.")
    add_screenshot(s16, "15_trang_lien_he.png", "Trang Liên hệ với bản đồ trụ sở 12 Trịnh Đình Thảo và thông tin hotline")

    prs.save(OUTPUT_PPTX)
    print(f"Presentation saved successfully to: {OUTPUT_PPTX}")


if __name__ == "__main__":
    build_presentation()
